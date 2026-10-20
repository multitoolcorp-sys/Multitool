import platform
import os
import sys
import shutil
import subprocess
from pathlib import Path
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, 
    QComboBox, QPushButton, QMessageBox, QFrame, QFileDialog, QProgressBar
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtGui import QFont, QCursor

def asegurar_site_packages_en_path():
    """Incorpora dinámicamente las carpetas site-packages y DLLs de Python local al sys.path del proceso."""
    base_progs = Path(os.path.expandvars(r"%LOCALAPPDATA%\Programs\Python"))
    if base_progs.exists():
        for py_dir in sorted(base_progs.glob("Python*"), reverse=True):
            sp = py_dir / "Lib" / "site-packages"
            if sp.exists():
                for sub in [sp, sp / "win32", sp / "win32" / "lib", sp / "Pythonwin"]:
                    if sub.exists() and str(sub) not in sys.path:
                        sys.path.append(str(sub))
                
                # DLLs de pywin32 (pythoncom y pywintypes)
                pyw32_sys = sp / "pywin32_system32"
                if pyw32_sys.exists():
                    os.environ["PATH"] = str(pyw32_sys) + os.pathsep + os.environ.get("PATH", "")
                    try:
                        os.add_dll_directory(str(pyw32_sys))
                    except Exception:
                        pass

asegurar_site_packages_en_path()

def obtener_python_sistema():
    """Localiza el ejecutable de Python del sistema con librerías completas."""
    if not getattr(sys, "frozen", False) and "multitool" not in Path(sys.executable).name.lower():
        return sys.executable

    base_progs = Path(os.path.expandvars(r"%LOCALAPPDATA%\Programs\Python"))
    if base_progs.exists():
        for py_dir in sorted(base_progs.glob("Python*"), reverse=True):
            cand = py_dir / "python.exe"
            if cand.exists():
                return str(cand)

    for cmd in ["python", "python3", "py"]:
        wh = shutil.which(cmd)
        if wh and "windowsapps" not in wh.lower() and "multitool" not in wh.lower():
            return wh

    return sys.executable

def ejecutar_script_en_python_sistema(codigo_script, in_file, out_file, timeout=120):
    """Ejecuta una rutina de conversión aislada usando el intérprete Python del sistema."""
    py_exe = obtener_python_sistema()
    if not py_exe or not Path(py_exe).exists():
        return False, "No se encontró el ejecutable de Python del sistema."
    try:
        startupinfo = None
        if os.name == 'nt':
            startupinfo = subprocess.STARTUPINFO()
            startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            startupinfo.wShowWindow = 0  # SW_HIDE
        res = subprocess.run(
            [py_exe, "-c", codigo_script, str(in_file), str(out_file)],
            capture_output=True,
            text=True,
            startupinfo=startupinfo,
            timeout=timeout
        )
        if res.returncode == 0 and Path(out_file).exists() and Path(out_file).stat().st_size > 0:
            return True, ""
        err = res.stderr.strip() or res.stdout.strip() or f"Código de salida: {res.returncode}"
        return False, err
    except Exception as e:
        return False, str(e)

class ConversionThread(QThread):
    progress = pyqtSignal(int)
    finished = pyqtSignal(str, bool)

    def __init__(self, in_file, out_file):
        super().__init__()
        self.in_file = Path(in_file)
        self.out_file = Path(out_file)

    def run(self):
        com_initialized = False
        try:
            import pythoncom
            pythoncom.CoInitialize()
            com_initialized = True
        except Exception:
            pass

        try:
            out_file = self.out_file
            ext_in = self.in_file.suffix.lower()
            out_ext = out_file.suffix.lower()
                
            self.progress.emit(20)
            
            # --- CONVERSIÓN DE IMÁGENES ---
            if ext_in in ['.jpg', '.jpeg', '.png', '.ico', '.webp', '.bmp', '.gif', '.tiff']:
                try:
                    from PIL import Image
                    img = Image.open(self.in_file)
                    
                    if out_ext == '.ico':
                        img.save(out_file, format='ICO', sizes=[(256, 256)])
                    elif out_ext in ['.jpg', '.jpeg']:
                        img.convert('RGB').save(out_file, format='JPEG', quality=95)
                    elif out_ext == '.pdf':
                        img.convert('RGB').save(out_file, format='PDF')
                    else:
                        img.save(out_file)
                        
                    self.progress.emit(100)
                    self.finished.emit(str(out_file), True)
                    return
                except Exception as e_img:
                    self.finished.emit(f"Error convirtiendo imagen: {e_img}", False)
                    return
            
            # --- CONVERSIÓN DE TEXTO A PDF ---
            if ext_in in ['.txt', '.md', '.csv'] and out_ext == '.pdf':
                convertido = False
                # Intento 1: reportlab
                try:
                    from reportlab.lib.pagesizes import letter
                    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                    from reportlab.lib.styles import getSampleStyleSheet
                    with open(self.in_file, 'r', encoding='utf-8', errors='replace') as f:
                        lines = f.readlines()
                    pdf_doc = SimpleDocTemplate(str(out_file), pagesize=letter)
                    styles = getSampleStyleSheet()
                    story = []
                    for line in lines[:800]:
                        if line.strip():
                            safe_txt = line.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                            story.append(Paragraph(safe_txt, styles['Normal']))
                            story.append(Spacer(1, 3))
                    pdf_doc.build(story)
                    convertido = True
                except Exception:
                    pass

                if not convertido:
                    try:
                        from fpdf import FPDF
                        with open(self.in_file, 'r', encoding='utf-8', errors='replace') as f:
                            text = f.read()
                        pdf = FPDF()
                        pdf.add_page()
                        pdf.set_font("Arial", size=10)
                        pdf.multi_cell(0, 5, text.encode('latin-1', 'replace').decode('latin-1'))
                        pdf.output(str(out_file))
                        convertido = True
                    except Exception as e_fpdf:
                        self.finished.emit(f"Error al generar PDF de texto: {e_fpdf}", False)
                        return

                if convertido and out_file.exists() and out_file.stat().st_size > 0:
                    self.progress.emit(100)
                    self.finished.emit(str(out_file), True)
                    return

            # --- CONVERSIÓN PDF A WORD (.pdf -> .docx) ---
            if ext_in == '.pdf' and out_ext == '.docx':
                convertido = False
                err_msg = ""
                # Intento 1: pdf2docx Converter (en proceso)
                try:
                    self.progress.emit(40)
                    from pdf2docx import Converter
                    cv = Converter(str(self.in_file))
                    cv.convert(str(out_file), start=0, end=None)
                    cv.close()
                    convertido = True
                except Exception as e_p2d:
                    err_msg = str(e_p2d)

                # Intento 2 (Fallback): PyMuPDF (fitz) + python-docx (en proceso)
                if not convertido:
                    try:
                        self.progress.emit(60)
                        import fitz
                        import docx
                        doc_fitz = fitz.open(str(self.in_file))
                        doc_word = docx.Document()
                        for page in doc_fitz:
                            text = page.get_text("text")
                            if text and text.strip():
                                for line in text.splitlines():
                                    if line.strip():
                                        doc_word.add_paragraph(line)
                            doc_word.add_page_break()
                        doc_fitz.close()
                        doc_word.save(str(out_file))
                        convertido = True
                    except Exception as e_fb:
                        err_msg += f" | Fallback PyMuPDF: {e_fb}"

                # Intento 3 (Fallback Subproceso Python del Sistema con librerías nativas)
                if not convertido:
                    self.progress.emit(75)
                    script_p2d = (
                        "import sys\n"
                        "from pdf2docx import Converter\n"
                        "cv = Converter(sys.argv[1])\n"
                        "cv.convert(sys.argv[2], start=0, end=None)\n"
                        "cv.close()\n"
                    )
                    ok, sub_err = ejecutar_script_en_python_sistema(script_p2d, self.in_file, out_file)
                    if ok:
                        convertido = True
                    else:
                        script_fitz = (
                            "import sys, fitz, docx\n"
                            "doc_fitz = fitz.open(sys.argv[1])\n"
                            "doc_word = docx.Document()\n"
                            "for page in doc_fitz:\n"
                            "    text = page.get_text('text')\n"
                            "    if text and text.strip():\n"
                            "        for line in text.splitlines():\n"
                            "            if line.strip():\n"
                            "                doc_word.add_paragraph(line)\n"
                            "    doc_word.add_page_break()\n"
                            "doc_fitz.close()\n"
                            "doc_word.save(sys.argv[2])\n"
                        )
                        ok_fb, sub_err_fb = ejecutar_script_en_python_sistema(script_fitz, self.in_file, out_file)
                        if ok_fb:
                            convertido = True
                        else:
                            err_msg += f" | Subproceso: {sub_err or sub_err_fb}"

                if convertido and out_file.exists() and out_file.stat().st_size > 0:
                    self.progress.emit(100)
                    self.finished.emit(str(out_file), True)
                    return
                else:
                    self.finished.emit(f"No se pudo convertir PDF a Word: {err_msg}", False)
                    return
                    
            # --- CONVERSIÓN WORD A PDF (.docx -> .pdf) ---
            if ext_in == '.docx' and out_ext == '.pdf':
                convertido = False
                err_msg = ""
                # Intento 1: Word COM nativo de Windows (fidelidad 100%)
                try:
                    self.progress.emit(40)
                    import win32com.client
                    word = win32com.client.DispatchEx("Word.Application")
                    word.Visible = False
                    word.DisplayAlerts = False
                    doc = word.Documents.Open(str(self.in_file.resolve()))
                    # wdFormatPDF = 17
                    doc.SaveAs(str(out_file.resolve()), FileFormat=17)
                    doc.Close(False)
                    word.Quit()
                    convertido = True
                except Exception as e_com:
                    err_msg = f"Word COM: {e_com}"

                # Intento 2: docx2pdf
                if not convertido:
                    try:
                        self.progress.emit(60)
                        from docx2pdf import convert
                        convert(str(self.in_file), str(out_file))
                        convertido = True
                    except Exception as e_d2p:
                        err_msg += f" | docx2pdf: {e_d2p}"

                # Intento 3 (Fallback Universal): python-docx + reportlab
                if not convertido:
                    try:
                        self.progress.emit(75)
                        import docx
                        from reportlab.lib.pagesizes import letter
                        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                        from reportlab.lib.styles import getSampleStyleSheet
                        
                        doc_in = docx.Document(str(self.in_file))
                        pdf_doc = SimpleDocTemplate(str(out_file), pagesize=letter)
                        styles = getSampleStyleSheet()
                        story = []
                        for p in doc_in.paragraphs:
                            if p.text and p.text.strip():
                                safe_txt = p.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                                story.append(Paragraph(safe_txt, styles['Normal']))
                                story.append(Spacer(1, 5))
                        if story:
                            pdf_doc.build(story)
                            convertido = True
                    except Exception as e_rl:
                        err_msg += f" | Fallback reportlab: {e_rl}"

                # Intento 4 (Fallback Subproceso Python del Sistema):
                if not convertido:
                    self.progress.emit(85)
                    script_w2p = (
                        "import sys\n"
                        "try:\n"
                        "    import win32com.client, pythoncom\n"
                        "    pythoncom.CoInitialize()\n"
                        "    word = win32com.client.DispatchEx('Word.Application')\n"
                        "    word.Visible = False\n"
                        "    word.DisplayAlerts = False\n"
                        "    doc = word.Documents.Open(sys.argv[1])\n"
                        "    doc.SaveAs(sys.argv[2], FileFormat=17)\n"
                        "    doc.Close(False)\n"
                        "    word.Quit()\n"
                        "except Exception:\n"
                        "    try:\n"
                        "        from docx2pdf import convert\n"
                        "        convert(sys.argv[1], sys.argv[2])\n"
                        "    except Exception:\n"
                        "        import docx\n"
                        "        from reportlab.lib.pagesizes import letter\n"
                        "        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer\n"
                        "        from reportlab.lib.styles import getSampleStyleSheet\n"
                        "        doc_in = docx.Document(sys.argv[1])\n"
                        "        pdf_doc = SimpleDocTemplate(sys.argv[2], pagesize=letter)\n"
                        "        styles = getSampleStyleSheet()\n"
                        "        story = []\n"
                        "        for p in doc_in.paragraphs:\n"
                        "            if p.text and p.text.strip():\n"
                        "                safe_txt = p.text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')\n"
                        "                story.append(Paragraph(safe_txt, styles['Normal']))\n"
                        "                story.append(Spacer(1, 5))\n"
                        "        if story:\n"
                        "            pdf_doc.build(story)\n"
                    )
                    ok_w, sub_w_err = ejecutar_script_en_python_sistema(script_w2p, self.in_file, out_file)
                    if ok_w:
                        convertido = True
                    else:
                        err_msg += f" | Subproceso Word: {sub_w_err}"

                if convertido and out_file.exists() and out_file.stat().st_size > 0:
                    self.progress.emit(100)
                    self.finished.emit(str(out_file), True)
                    return
                else:
                    self.finished.emit(f"Error al convertir Word a PDF: {err_msg}", False)
                    return

            # --- CONVERSIÓN DE HOJAS DE CÁLCULO EXCEL (.xlsx / .xls) ---
            if ext_in in ['.xlsx', '.xls']:
                # 1. XLSX a CSV
                if out_ext == '.csv':
                    try:
                        import openpyxl
                        import csv
                        wb = openpyxl.load_workbook(self.in_file, data_only=True)
                        ws = wb.active
                        with open(out_file, 'w', newline='', encoding='utf-8-sig') as f:
                            writer = csv.writer(f)
                            for row in ws.iter_rows(values_only=True):
                                if any(cell is not None for cell in row):
                                    writer.writerow([str(c) if c is not None else "" for c in row])
                        self.progress.emit(100)
                        self.finished.emit(str(out_file), True)
                        return
                    except Exception as e_csv:
                        script_csv = (
                            "import sys, openpyxl, csv\n"
                            "wb = openpyxl.load_workbook(sys.argv[1], data_only=True)\n"
                            "ws = wb.active\n"
                            "with open(sys.argv[2], 'w', newline='', encoding='utf-8-sig') as f:\n"
                            "    writer = csv.writer(f)\n"
                            "    for row in ws.iter_rows(values_only=True):\n"
                            "        if any(cell is not None for cell in row):\n"
                            "            writer.writerow([str(c) if c is not None else '' for c in row])\n"
                        )
                        ok, sub_e = ejecutar_script_en_python_sistema(script_csv, self.in_file, out_file)
                        if ok:
                            self.progress.emit(100)
                            self.finished.emit(str(out_file), True)
                            return
                        self.finished.emit(f"Error al convertir Excel a CSV: {e_csv} | Subproceso: {sub_e}", False)
                        return

                # 2. XLSX a JSON
                if out_ext == '.json':
                    try:
                        import openpyxl
                        import json
                        wb = openpyxl.load_workbook(self.in_file, data_only=True)
                        ws = wb.active
                        rows = list(ws.iter_rows(values_only=True))
                        if not rows:
                            raise ValueError("El archivo Excel no contiene datos.")
                        
                        headers = [str(c or f"col_{idx+1}").strip() for idx, c in enumerate(rows[0])]
                        data = []
                        for row in rows[1:]:
                            if any(c is not None for c in row):
                                item = {}
                                for idx, val in enumerate(row):
                                    col_name = headers[idx] if idx < len(headers) else f"col_{idx+1}"
                                    item[col_name] = str(val) if val is not None else None
                                data.append(item)
                                
                        with open(out_file, 'w', encoding='utf-8') as f:
                            json.dump(data, f, ensure_ascii=False, indent=2)
                        self.progress.emit(100)
                        self.finished.emit(str(out_file), True)
                        return
                    except Exception as e_json:
                        script_json = (
                            "import sys, openpyxl, json\n"
                            "wb = openpyxl.load_workbook(sys.argv[1], data_only=True)\n"
                            "ws = wb.active\n"
                            "rows = list(ws.iter_rows(values_only=True))\n"
                            "if rows:\n"
                            "    headers = [str(c or f'col_{idx+1}').strip() for idx, c in enumerate(rows[0])]\n"
                            "    data = []\n"
                            "    for row in rows[1:]:\n"
                            "        if any(c is not None for c in row):\n"
                            "            item = {headers[idx] if idx < len(headers) else f'col_{idx+1}': (str(val) if val is not None else None) for idx, val in enumerate(row)}\n"
                            "            data.append(item)\n"
                            "    with open(sys.argv[2], 'w', encoding='utf-8') as f:\n"
                            "        json.dump(data, f, ensure_ascii=False, indent=2)\n"
                        )
                        ok, sub_e = ejecutar_script_en_python_sistema(script_json, self.in_file, out_file)
                        if ok:
                            self.progress.emit(100)
                            self.finished.emit(str(out_file), True)
                            return
                        self.finished.emit(f"Error al convertir Excel a JSON: {e_json} | Subproceso: {sub_e}", False)
                        return

                # 3. XLSX a TXT
                if out_ext == '.txt':
                    try:
                        import openpyxl
                        wb = openpyxl.load_workbook(self.in_file, data_only=True)
                        ws = wb.active
                        with open(out_file, 'w', encoding='utf-8') as f:
                            for row in ws.iter_rows(values_only=True):
                                if any(cell is not None for cell in row):
                                    line = "\t".join(str(c) if c is not None else "" for c in row)
                                    f.write(line + "\n")
                        self.progress.emit(100)
                        self.finished.emit(str(out_file), True)
                        return
                    except Exception as e_txt:
                        script_txt = (
                            "import sys, openpyxl\n"
                            "wb = openpyxl.load_workbook(sys.argv[1], data_only=True)\n"
                            "ws = wb.active\n"
                            "with open(sys.argv[2], 'w', encoding='utf-8') as f:\n"
                            "    for row in ws.iter_rows(values_only=True):\n"
                            "        if any(cell is not None for cell in row):\n"
                            "            f.write('\\t'.join(str(c) if c is not None else '' for c in row) + '\\n')\n"
                        )
                        ok, sub_e = ejecutar_script_en_python_sistema(script_txt, self.in_file, out_file)
                        if ok:
                            self.progress.emit(100)
                            self.finished.emit(str(out_file), True)
                            return
                        self.finished.emit(f"Error al convertir Excel a TXT: {e_txt} | Subproceso: {sub_e}", False)
                        return

                # 4. XLSX a HTML
                if out_ext == '.html':
                    try:
                        import openpyxl
                        wb = openpyxl.load_workbook(self.in_file, data_only=True)
                        ws = wb.active
                        html_parts = [
                            "<!DOCTYPE html><html><head><meta charset='utf-8'><title>Excel Export</title>",
                            "<style>table {border-collapse: collapse; width: 100%; font-family: Segoe UI, sans-serif; font-size: 13px;} th, td {border: 1px solid #ddd; padding: 8px;} tr:nth-child(even){background-color: #f9f9f9;} th {background-color: #6366F1; color: white; text-align: left;}</style>",
                            "</head><body><table>"
                        ]
                        rows = list(ws.iter_rows(values_only=True))
                        if rows:
                            html_parts.append("<thead><tr>" + "".join(f"<th>{c or ''}</th>" for c in rows[0]) + "</tr></thead><tbody>")
                            for r in rows[1:]:
                                if any(c is not None for c in r):
                                    html_parts.append("<tr>" + "".join(f"<td>{c if c is not None else ''}</td>" for c in r) + "</tr>")
                            html_parts.append("</tbody>")
                        html_parts.append("</table></body></html>")
                        with open(out_file, 'w', encoding='utf-8') as f:
                            f.write("\n".join(html_parts))
                        self.progress.emit(100)
                        self.finished.emit(str(out_file), True)
                        return
                    except Exception as e_html:
                        script_html = (
                            "import sys, openpyxl\n"
                            "wb = openpyxl.load_workbook(sys.argv[1], data_only=True)\n"
                            "ws = wb.active\n"
                            "rows = list(ws.iter_rows(values_only=True))\n"
                            "parts = ['<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>Excel</title>', '<style>table{border-collapse:collapse;width:100%;font-family:sans-serif;}th,td{border:1px solid #ddd;padding:8px;}th{background:#6366F1;color:white;}</style></head><body><table>']\n"
                            "if rows:\n"
                            "    parts.append('<thead><tr>' + ''.join(f'<th>{c or \"\"}</th>' for c in rows[0]) + '</tr></thead><tbody>')\n"
                            "    for r in rows[1:]:\n"
                            "        if any(c is not None for c in r):\n"
                            "            parts.append('<tr>' + ''.join(f'<td>{c if c is not None else \"\"}</td>' for c in r) + '</tr>')\n"
                            "    parts.append('</tbody>')\n"
                            "parts.append('</table></body></html>')\n"
                            "with open(sys.argv[2], 'w', encoding='utf-8') as f:\n"
                            "    f.write('\\n'.join(parts))\n"
                        )
                        ok, sub_e = ejecutar_script_en_python_sistema(script_html, self.in_file, out_file)
                        if ok:
                            self.progress.emit(100)
                            self.finished.emit(str(out_file), True)
                            return
                        self.finished.emit(f"Error al convertir Excel a HTML: {e_html} | Subproceso: {sub_e}", False)
                        return

                # 5. XLSX a PDF
                if out_ext == '.pdf':
                    convertido = False
                    err_msg = ""
                    # Intento 1: Excel COM nativo de Windows (si Microsoft Excel está instalado)
                    try:
                        self.progress.emit(40)
                        import win32com.client
                        excel = win32com.client.DispatchEx("Excel.Application")
                        excel.Visible = False
                        excel.DisplayAlerts = False
                        wb = excel.Workbooks.Open(str(self.in_file.resolve()))
                        wb.ExportAsFixedFormat(0, str(out_file.resolve()))
                        wb.Close(False)
                        excel.Quit()
                        convertido = True
                    except Exception as e_excom:
                        err_msg = f"Excel COM: {e_excom}"

                    # Intento 2: openpyxl + reportlab (Motor universal sin necesidad de Excel)
                    if not convertido:
                        try:
                            self.progress.emit(60)
                            import openpyxl
                            from reportlab.lib.pagesizes import letter, landscape
                            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
                            from reportlab.lib.styles import getSampleStyleSheet
                            from reportlab.lib import colors

                            wb = openpyxl.load_workbook(self.in_file, data_only=True)
                            ws = wb.active
                            rows = list(ws.iter_rows(values_only=True))
                            
                            table_data = []
                            for row in rows:
                                if any(c is not None for c in row):
                                    table_data.append([str(c) if c is not None else "" for c in row])
                            
                            if table_data:
                                max_cols = max(len(r) for r in table_data)
                                for r in table_data:
                                    while len(r) < max_cols:
                                        r.append("")
                                        
                                psize = landscape(letter) if max_cols > 5 else letter
                                pdf_doc = SimpleDocTemplate(str(out_file), pagesize=psize, leftMargin=18, rightMargin=18, topMargin=20, bottomMargin=20)
                                styles = getSampleStyleSheet()
                                story = [Paragraph(f"<b>{self.in_file.stem}</b>", styles['Title']), Spacer(1, 10)]
                                
                                formatted_data = []
                                for r_idx, row in enumerate(table_data[:200]):
                                    row_cells = []
                                    for c in row:
                                        style = styles['Normal']
                                        style.fontSize = 8
                                        style.leading = 10
                                        txt = str(c)[:60].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                                        row_cells.append(Paragraph(txt, style))
                                    formatted_data.append(row_cells)
                                    
                                t = Table(formatted_data, repeatRows=1)
                                t.setStyle(TableStyle([
                                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366F1')),
                                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
                                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')])
                                ]))
                                story.append(t)
                                pdf_doc.build(story)
                                convertido = True
                        except Exception as e_rl:
                            err_msg += f" | reportlab: {e_rl}"

                    # Intento 3 (Fallback Subproceso Python del Sistema):
                    if not convertido:
                        self.progress.emit(85)
                        script_xl = (
                            "import sys\n"
                            "try:\n"
                            "    import win32com.client, pythoncom\n"
                            "    pythoncom.CoInitialize()\n"
                            "    excel = win32com.client.DispatchEx('Excel.Application')\n"
                            "    excel.Visible = False\n"
                            "    excel.DisplayAlerts = False\n"
                            "    wb = excel.Workbooks.Open(sys.argv[1])\n"
                            "    wb.ExportAsFixedFormat(0, sys.argv[2])\n"
                            "    wb.Close(False)\n"
                            "    excel.Quit()\n"
                            "except Exception:\n"
                            "    import openpyxl\n"
                            "    from reportlab.lib.pagesizes import letter, landscape\n"
                            "    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer\n"
                            "    from reportlab.lib.styles import getSampleStyleSheet\n"
                            "    from reportlab.lib import colors\n"
                            "    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)\n"
                            "    ws = wb.active\n"
                            "    rows = list(ws.iter_rows(values_only=True))\n"
                            "    table_data = []\n"
                            "    for row in rows:\n"
                            "        if any(c is not None for c in row):\n"
                            "            table_data.append([str(c) if c is not None else '' for c in row])\n"
                            "    if table_data:\n"
                            "        max_cols = max(len(r) for r in table_data)\n"
                            "        for r in table_data:\n"
                            "            while len(r) < max_cols:\n"
                            "                r.append('')\n"
                            "        psize = landscape(letter) if max_cols > 5 else letter\n"
                            "        pdf_doc = SimpleDocTemplate(sys.argv[2], pagesize=psize, leftMargin=18, rightMargin=18, topMargin=20, bottomMargin=20)\n"
                            "        styles = getSampleStyleSheet()\n"
                            "        story = [Paragraph('<b>Excel Export</b>', styles['Title']), Spacer(1, 10)]\n"
                            "        formatted_data = []\n"
                            "        for row in table_data[:200]:\n"
                            "            row_cells = []\n"
                            "            for c in row:\n"
                            "                style = styles['Normal']\n"
                            "                style.fontSize = 8\n"
                            "                txt = str(c)[:60].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')\n"
                            "                row_cells.append(Paragraph(txt, style))\n"
                            "            formatted_data.append(row_cells)\n"
                            "        t = Table(formatted_data, repeatRows=1)\n"
                            "        t.setStyle(TableStyle([\n"
                            "            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366F1')),\n"
                            "            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),\n"
                            "            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC'))\n"
                            "        ]))\n"
                            "        story.append(t)\n"
                            "        pdf_doc.build(story)\n"
                        )
                        ok_xl, sub_xl_err = ejecutar_script_en_python_sistema(script_xl, self.in_file, out_file)
                        if ok_xl:
                            convertido = True
                        else:
                            err_msg += f" | Subproceso Excel: {sub_xl_err}"

                    if convertido and out_file.exists() and out_file.stat().st_size > 0:
                        self.progress.emit(100)
                        self.finished.emit(str(out_file), True)
                        return
                    else:
                        self.finished.emit(f"Error convirtiendo Excel a PDF: {err_msg}", False)
                        return

            # --- CONVERSIÓN CSV / JSON A EXCEL (.xlsx) ---
            if ext_in == '.csv' and out_ext == '.xlsx':
                try:
                    import csv
                    import openpyxl
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    with open(self.in_file, 'r', encoding='utf-8', errors='replace') as f:
                        reader = csv.reader(f)
                        for r_idx, row in enumerate(reader, start=1):
                            for c_idx, val in enumerate(row, start=1):
                                ws.cell(row=r_idx, column=c_idx, value=val)
                    wb.save(out_file)
                    self.progress.emit(100)
                    self.finished.emit(str(out_file), True)
                    return
                except Exception as e_c2x:
                    script_c2x = (
                        "import sys, csv, openpyxl\n"
                        "wb = openpyxl.Workbook()\n"
                        "ws = wb.active\n"
                        "with open(sys.argv[1], 'r', encoding='utf-8', errors='replace') as f:\n"
                        "    reader = csv.reader(f)\n"
                        "    for r_idx, row in enumerate(reader, start=1):\n"
                        "        for c_idx, val in enumerate(row, start=1):\n"
                        "            ws.cell(row=r_idx, column=c_idx, value=val)\n"
                        "wb.save(sys.argv[2])\n"
                    )
                    ok, sub_e = ejecutar_script_en_python_sistema(script_c2x, self.in_file, out_file)
                    if ok:
                        self.progress.emit(100)
                        self.finished.emit(str(out_file), True)
                        return
                    self.finished.emit(f"Error convirtiendo CSV a Excel: {e_c2x} | Subproceso: {sub_e}", False)
                    return

            if ext_in == '.json' and out_ext == '.xlsx':
                try:
                    import json
                    import openpyxl
                    with open(self.in_file, 'r', encoding='utf-8') as f:
                        data = json.load(f)
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    if isinstance(data, list) and data and isinstance(data[0], dict):
                        headers = list(data[0].keys())
                        ws.append(headers)
                        for item in data:
                            if isinstance(item, dict):
                                ws.append([item.get(h, "") for h in headers])
                    elif isinstance(data, list):
                        for item in data:
                            ws.append(item if isinstance(item, list) else [item])
                    wb.save(out_file)
                    self.progress.emit(100)
                    self.finished.emit(str(out_file), True)
                    return
                except Exception as e_j2x:
                    script_j2x = (
                        "import sys, json, openpyxl\n"
                        "with open(sys.argv[1], 'r', encoding='utf-8') as f:\n"
                        "    data = json.load(f)\n"
                        "wb = openpyxl.Workbook()\n"
                        "ws = wb.active\n"
                        "if isinstance(data, list) and data and isinstance(data[0], dict):\n"
                        "    headers = list(data[0].keys())\n"
                        "    ws.append(headers)\n"
                        "    for item in data:\n"
                        "        if isinstance(item, dict):\n"
                        "            ws.append([item.get(h, '') for h in headers])\n"
                        "elif isinstance(data, list):\n"
                        "    for item in data:\n"
                        "        ws.append(item if isinstance(item, list) else [item])\n"
                        "wb.save(sys.argv[2])\n"
                    )
                    ok, sub_e = ejecutar_script_en_python_sistema(script_j2x, self.in_file, out_file)
                    if ok:
                        self.progress.emit(100)
                        self.finished.emit(str(out_file), True)
                        return
                    self.finished.emit(f"Error convirtiendo JSON a Excel: {e_j2x} | Subproceso: {sub_e}", False)
                    return
                    
            # --- CONVERSIÓN PDF/WORD A IMÁGENES (.png / .jpg) ---
            if (ext_in in ['.pdf', '.docx']) and (out_ext in ['.png', '.jpg', '.jpeg']):
                try:
                    import tempfile
                    import uuid
                    import zipfile
                    
                    pdf_to_read = self.in_file
                    temp_pdf = None
                    
                    if ext_in == '.docx':
                        # Convertir Word a PDF temporal primero
                        try:
                            temp_pdf = Path(tempfile.gettempdir()) / f"temp_{uuid.uuid4().hex}.pdf"
                            from docx2pdf import convert
                            convert(str(self.in_file), str(temp_pdf))
                            pdf_to_read = temp_pdf
                        except Exception:
                            pass

                    # Abrir PDF y extraer páginas
                    try:
                        import fitz
                        doc = fitz.open(pdf_to_read)
                        page_count = len(doc)
                        if page_count == 0:
                            raise ValueError("El documento no contiene páginas.")
                    except Exception as e_fitz:
                        self.finished.emit(f"Error al abrir documento para extracción de imágenes: {e_fitz}", False)
                        return

                    if page_count == 1:
                        page = doc.load_page(0)
                        pix = page.get_pixmap(dpi=150)
                        pix.save(str(out_file))
                        final_out = out_file
                    else:
                        zip_path = out_file.with_suffix('.zip')
                        with zipfile.ZipFile(zip_path, 'w') as zip_file:
                            for page_num in range(page_count):
                                page = doc.load_page(page_num)
                                pix = page.get_pixmap(dpi=150)
                                img_temp = Path(tempfile.gettempdir()) / f"pagina_{page_num + 1}{out_ext}"
                                pix.save(str(img_temp))
                                zip_file.write(img_temp, arcname=f"pagina_{page_num + 1}{out_ext}")
                                try: img_temp.unlink()
                                except: pass
                        final_out = zip_path

                    doc.close()
                    if temp_pdf and temp_pdf.exists():
                        try: temp_pdf.unlink()
                        except: pass

                    self.progress.emit(100)
                    self.finished.emit(str(final_out), True)
                    return
                except Exception as e_img:
                    self.finished.emit(f"Error convirtiendo a imágenes: {str(e_img)}", False)
                    return
                    
            # --- CONVERSIÓN DE TEXTO/CÓDIGO ---
            if ext_in in ['.txt', '.md', '.html', '.json', '.csv', '.xml'] and out_ext in ['.txt', '.md', '.html', '.json', '.csv', '.xml']:
                with open(self.in_file, 'r', encoding='utf-8', errors='ignore') as f:
                    data = f.read()
                with open(out_file, 'w', encoding='utf-8') as f:
                    f.write(data)
                self.progress.emit(100)
                self.finished.emit(str(out_file), True)
                return

            self.finished.emit(f"No hay motores locales disponibles para convertir {ext_in} a {out_ext}.", False)
            
        except Exception as e:
            self.finished.emit(f"Error procesando conversión: {str(e)}", False)
        finally:
            if com_initialized:
                try:
                    import pythoncom
                    pythoncom.CoUninitialize()
                except Exception:
                    pass


